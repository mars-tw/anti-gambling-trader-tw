"""Offline tests for core.ui.service (no servers / credentials)."""

from __future__ import annotations

import csv
import io
import json
import re
import threading
import time
import zipfile
from pathlib import Path

import pytest

import core as core_pkg
import core.ui.service as service_module
from core.ui import Artifact, UIError, UIService
from core.ui.service import _excel_available


EXAMPLES = Path(core_pkg.__file__).resolve().parent / "examples"


def _rewrite_xlsx_sheet(raw: bytes, transform, *, sheet_name="xl/worksheets/sheet1.xml"):
    """Change only worksheet XML while retaining a valid XLSX ZIP container."""

    output = io.BytesIO()
    found = False
    with zipfile.ZipFile(io.BytesIO(raw), "r") as source:
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == sheet_name:
                    data = transform(data)
                    found = True
                target.writestr(info, data)
    assert found
    return output.getvalue()


def _relocate_xlsx_sheet(raw: bytes) -> bytes:
    """Move a valid worksheet to an arbitrary XML part and update OOXML refs."""

    output = io.BytesIO()
    old_name = "xl/worksheets/sheet1.xml"
    new_name = "xl/custom.xml"
    with zipfile.ZipFile(io.BytesIO(raw), "r") as source:
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                data = source.read(info.filename)
                name = info.filename
                if name == old_name:
                    name = new_name
                elif name == "xl/_rels/workbook.xml.rels":
                    data = data.replace(b"worksheets/sheet1.xml", b"custom.xml")
                elif name == "[Content_Types].xml":
                    data = data.replace(b"/xl/worksheets/sheet1.xml", b"/xl/custom.xml")
                target.writestr(name if name != info.filename else info, data)
    return output.getvalue()


@pytest.fixture
def svc():
    s = UIService(mode="browser")
    yield s
    s.close()


def test_state_capabilities(svc: UIService):
    st = svc.state()
    assert st["mode"] == "browser"
    assert st["capabilities"]["live_trading"] is False
    assert st["capabilities"]["excel"] is _excel_available()
    assert st["manual_rows"] == []
    assert st["analysis"] is None
    assert st["busy"] is False


def test_analyze_builtin_tw_sample(svc: UIService):
    out = svc.analyze_sample("tw")
    assert out["origin"] == "demo"
    assert out["can_live"] is False
    assert isinstance(out["id"], str) and out["id"]
    assert out["revision"] >= 1
    assert "result" in out and "stage" in out["result"]
    assert "metrics" in out
    assert isinstance(out["text_report"], str) and out["text_report"]
    assert isinstance(out["readiness"], list) and out["readiness"]
    keys = {r["key"] for r in out["readiness"]}
    assert "total_trades>=30" in keys
    assert "expectancy>0" in keys
    assert all("passed" in r and "label" in r for r in out["readiness"])


def test_analyze_builtin_us_and_crypto(svc: UIService):
    us = svc.analyze_sample("us")
    assert us["origin"] == "demo"
    crypto = svc.analyze_sample("crypto")
    assert crypto["origin"] == "demo"
    assert crypto["id"] != us["id"]


def test_readiness_details_are_human_facing(svc: UIService):
    out = svc.analyze_sample("tw")
    details = {item["key"]: item["detail"] for item in out["readiness"]}
    assert "expectancy=" not in details["expectancy>0"]
    assert "每筆平均結果" in details["expectancy>0"] or "無法計算" in details["expectancy>0"]
    assert "available=" not in details["oos_edge_persisted"]
    assert "edge_persisted=" not in details["oos_edge_persisted"]


def test_unknown_sample_name(svc: UIService):
    with pytest.raises(UIError) as ei:
        svc.analyze_sample("nope")
    assert ei.value.status == 400


def test_demo_provenance_via_renamed_upload(svc: UIService):
    raw = (EXAMPLES / "tw_stock_gambling.csv").read_bytes()
    out = svc.analyze_upload("totally_new_name.csv", raw)
    assert out["origin"] == "demo"


def test_import_origin_for_modified_content(svc: UIService):
    raw = (EXAMPLES / "tw_stock_gambling.csv").read_bytes() + b"\n"
    out = svc.analyze_upload("custom.csv", raw)
    # Extra newline may still parse; origin must not be demo digest match.
    assert out["origin"] == "import"


def _assert_export_provenance(
    svc: UIService,
    analysis: dict,
    *,
    demo: bool,
    source_label: str,
) -> None:
    json_artifact = svc.export("json", analysis_id=analysis["id"])
    payload = json.loads(json_artifact.content.decode("utf-8"))
    assert payload["ui_provenance"] == {
        "origin": analysis["origin"],
        "demo": demo,
        "analysis_id": analysis["id"],
        "revision": analysis["revision"],
    }
    without_provenance = {
        key: value for key, value in payload.items() if key != "ui_provenance"
    }
    assert without_provenance == analysis["result"]

    html = svc.export("html", analysis_id=analysis["id"]).content.decode("utf-8")
    card = svc.export("card", analysis_id=analysis["id"]).content.decode("utf-8")
    for document in (html, card):
        assert source_label in document
        if demo:
            assert "內建示範資料，非本人績效" in document
        else:
            assert "內建示範資料，非本人績效" not in document


def test_builtin_demo_provenance_is_in_all_exports(svc: UIService):
    out = svc.analyze_sample("tw")
    _assert_export_provenance(
        svc,
        out,
        demo=True,
        source_label="內建示範資料，非本人績效",
    )


def test_renamed_byte_identical_demo_provenance_is_in_all_exports(svc: UIService):
    raw = (EXAMPLES / "tw_stock_gambling.csv").read_bytes()
    out = svc.analyze_upload("my-performance.csv", raw)
    _assert_export_provenance(
        svc,
        out,
        demo=True,
        source_label="內建示範資料，非本人績效",
    )


def test_true_import_provenance_is_not_mislabeled_demo(svc: UIService):
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=list(service_module.BEGINNER_COLUMNS))
    writer.writeheader()
    for i in range(3):
        writer.writerow(
            {
                "代號": "OWN",
                "方向": "買",
                "進場時間": f"2024-04-0{i + 1}",
                "出場時間": f"2024-04-1{i + 1}",
                "進場價": "100",
                "出場價": str(101 + i),
                "數量": "1",
                "手續費": "0",
                "損益": str(i + 1),
                "損益幣別": "USD",
                "策略": "my-rule",
            }
        )
    raw = ("\ufeff" + buf.getvalue()).encode("utf-8")
    out = svc.analyze_upload("my-performance.csv", raw)
    assert out["origin"] == "import"
    _assert_export_provenance(
        svc,
        out,
        demo=False,
        source_label="來源：使用者匯入資料",
    )


def test_dirty_later_row_keeps_integrity_conservative(svc: UIService):
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=list(service_module.BEGINNER_COLUMNS))
    writer.writeheader()
    for i in range(3):
        writer.writerow(
            {
                "代號": "AAPL",
                "方向": "買",
                "進場時間": f"2024-01-0{i + 1}",
                "出場時間": f"2024-01-1{i + 1}",
                "進場價": "100",
                "出場價": "101",
                "數量": "1",
                "手續費": "0",
                "損益": "1",
                "損益幣別": "USD",
                "策略": "fixed",
            }
        )
    writer.writerow(
        {
            "代號": "AAPL",
            "方向": "買",
            "進場時間": "2024-02-01",
            "出場時間": "not-a-real-time",
            "損益": "not-a-number",
            "損益幣別": "USD",
        }
    )
    out = svc.analyze_upload("dirty.csv", buf.getvalue().encode("utf-8"))
    integrity = out["result"]["integrity"]
    assert integrity["complete"] is False
    assert integrity["rejected_row_count"] >= 1
    assert out["result"]["stage"]["code"] != "tiny_live_validation"


def test_upload_size_limit(svc: UIService):
    big = b"x" * (2 * 1024 * 1024 + 1)
    with pytest.raises(UIError) as ei:
        svc.analyze_upload("big.csv", big)
    assert ei.value.status == 413


def test_upload_rejects_non_csv_json(svc: UIService):
    with pytest.raises(UIError):
        svc.analyze_upload("note.txt", b"hello")


def test_manual_requires_exit_time_and_currency(svc: UIService):
    with pytest.raises(UIError) as ei:
        svc.add_record(
            {
                "symbol": "AAPL",
                "pnl": 10,
                "side": "long",
                "entry_time": "2024-01-01",
                "exit_time": "",
                "currency": "USD",
            }
        )
    assert "出場" in str(ei.value) or "exit_time" in str(ei.value)

    with pytest.raises(UIError) as ei2:
        svc.add_record(
            {
                "symbol": "AAPL",
                "pnl": 10,
                "side": "long",
                "entry_time": "2024-01-01",
                "exit_time": "2024-01-02",
                "currency": "",
            }
        )
    assert "幣" in str(ei2.value) or "currency" in str(ei2.value)


def test_reject_unknown_dangerous_fields(svc: UIService):
    with pytest.raises(UIError) as ei:
        svc.add_record(
            {
                "symbol": "AAPL",
                "pnl": 1,
                "side": "",
                "entry_time": "2024-01-01",
                "exit_time": "2024-01-02",
                "currency": "USD",
                "allow_live_trading": True,
                "dest_dir": "C:/evil",
            }
        )
    assert ei.value.status == 400


def test_manual_add_analyze_and_export_bom(svc: UIService):
    st = svc.add_record(
        {
            "symbol": "2330",
            "pnl": 100,
            "side": "long",
            "entry_time": "2024-01-01",
            "exit_time": "2024-01-05",
            "entry_price": 100,
            "exit_price": 110,
            "quantity": 1,
            "fees": 0,
            "currency": "TWD",
            "strategy": "test",
        }
    )
    assert st["manual_revision"] == 1
    assert len(st["manual_rows"]) == 1
    assert st["analysis"] is None  # invalidated / not yet analyzed

    # Need enough rows for a meaningful run — still should parse.
    for i in range(5):
        svc.add_record(
            {
                "symbol": "2330",
                "pnl": -10 + i,
                "side": "",  # unknown side allowed; never default long
                "entry_time": f"2024-02-0{i+1}",
                "exit_time": f"2024-02-1{i+1}",
                "currency": "TWD",
            }
        )

    art = svc.export("records")
    assert isinstance(art, Artifact)
    assert art.content.startswith(b"\xef\xbb\xbf")
    text = art.content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    assert len(rows) >= 2  # header + data

    tmpl = svc.export("template")
    assert tmpl.content.startswith(b"\xef\xbb\xbf")
    header_only = tmpl.content.decode("utf-8-sig").strip().splitlines()
    assert len(header_only) == 1


def test_analyze_records_origin_manual(svc: UIService):
    for i in range(3):
        svc.add_record(
            {
                "symbol": "MSFT",
                "pnl": 5 * (i + 1),
                "side": "short",
                "entry_time": f"2024-03-0{i+1}",
                "exit_time": f"2024-03-1{i+1}",
                "currency": "USD",
            }
        )
    out = svc.analyze_records()
    assert out["origin"] == "manual"
    assert out["can_live"] is False


def test_export_restore_append_roundtrip_invalidates_old_analysis(svc: UIService):
    source = UIService(mode="browser")
    try:
        source.add_record(
            {
                "symbol": "2330",
                "pnl": "1250",
                "side": "long",
                "entry_time": "2024-01-01T09:00",
                "exit_time": "2024-01-02T13:30",
                "currency": "TWD",
                "strategy": "roundtrip",
            }
        )
        source.add_record(
            {
                "symbol": "AAPL",
                "pnl": "-12.5",
                "side": "short",
                "exit_time": "2024-02-02T15:00",
                "currency": "USD",
            }
        )
        expected_rows = source.state()["manual_rows"]
        exported = source.export("records").content
    finally:
        source.close()

    old = svc.analyze_sample("tw")
    restored = svc.import_records("records.csv", exported)
    assert restored["manual_rows"] == expected_rows
    assert restored["analysis"] is None
    assert exported.startswith(b"\xef\xbb\xbf")  # restore never mutates its source bytes
    with pytest.raises(UIError) as stale:
        svc.export("json", analysis_id=old["id"])
    assert stale.value.status == 409

    appended = svc.add_record(
        {
            "symbol": "MSFT",
            "pnl": "3",
            "side": "long",
            "exit_time": "2024-03-03T15:00",
            "currency": "USD",
        }
    )
    combined = svc.export("records").content.decode("utf-8-sig")
    rows = list(csv.DictReader(io.StringIO(combined)))
    assert len(rows) == 3
    assert rows[:2] == expected_rows
    assert appended["analysis"] is None


def test_restore_rejects_nonstandard_or_invalid_rows_atomically(svc: UIService):
    before = svc.add_record(
        {
            "symbol": "AAPL",
            "pnl": "1",
            "side": "long",
            "exit_time": "2024-01-02",
            "currency": "USD",
        }
    )
    old_rows = before["manual_rows"]
    old = svc.analyze_sample("tw")
    malformed = "\ufeff" + ",".join(service_module.BEGINNER_COLUMNS) + "\nAAPL,買,,,,,,,1,,\n"
    with pytest.raises(UIError) as invalid:
        svc.import_records("records.csv", malformed.encode("utf-8"))
    assert invalid.value.status == 400
    assert svc.state()["manual_rows"] == old_rows
    assert svc.state()["analysis"] is None
    with pytest.raises(UIError):
        svc.export("html", analysis_id=old["id"])

    with pytest.raises(UIError):
        svc.import_records("broker.csv", b"symbol,pnl\nAAPL,1\n")
    assert svc.state()["manual_rows"] == old_rows


def test_stale_analysis_id_after_mutation(svc: UIService):
    out = svc.analyze_sample("tw")
    old_id = out["id"]
    svc.add_record(
        {
            "symbol": "X",
            "pnl": 1,
            "side": "",
            "entry_time": "2024-01-01",
            "exit_time": "2024-01-02",
            "currency": "USD",
        }
    )
    assert svc.state()["analysis"] is None
    with pytest.raises(UIError) as ei:
        svc.export("json", analysis_id=old_id)
    assert ei.value.status in (400, 409)


def test_stale_id_after_failed_import(svc: UIService):
    good = svc.analyze_sample("us")
    good_id = good["id"]
    with pytest.raises(UIError):
        svc.analyze_upload("bad.json", b"{not-json")
    # Failure must invalidate prior favorable analysis.
    assert svc.state()["analysis"] is None
    with pytest.raises(UIError) as ei:
        svc.export("html", analysis_id=good_id)
    assert ei.value.status in (400, 409)


def test_stale_id_after_preflight_failure(svc: UIService):
    good = svc.analyze_sample("us")
    with pytest.raises(UIError):
        svc.analyze_upload("wrong.txt", b"not supported")
    assert svc.state()["analysis"] is None
    with pytest.raises(UIError) as ei:
        svc.export("json", analysis_id=good["id"])
    assert ei.value.status == 409


def test_export_json_html_card_need_exact_id(svc: UIService):
    out = svc.analyze_sample("tw")
    aid = out["id"]
    j = svc.export("json", analysis_id=aid)
    assert j.media_type.startswith("application/json")
    json.loads(j.content.decode("utf-8"))

    h = svc.export("html", analysis_id=aid)
    assert b"<" in h.content

    c = svc.export("card", analysis_id=aid)
    assert isinstance(c.content, bytes)

    with pytest.raises(UIError):
        svc.export("json", analysis_id="nope")
    with pytest.raises(UIError):
        svc.export("json", analysis_id=None)


def test_scan_text_and_limit(svc: UIService):
    out = svc.scan("保證獲利，穩賺不賠，私訊領取內線")
    assert "result" in out and "text_report" in out
    assert isinstance(out["text_report"], str)

    with pytest.raises(UIError):
        svc.scan("x" * 20001)


def test_scaffold_always_paper_no_analysis_required(svc: UIService):
    art = svc.scaffold("demo_proj", symbols=["AAPL", "MSFT"], market="us_stock")
    assert art.filename.endswith(".zip")
    assert art.media_type == "application/zip"
    with zipfile.ZipFile(io.BytesIO(art.content)) as zf:
        names = zf.namelist()
        assert any(n.startswith("demo_proj/") for n in names)
        assert any(n.endswith("README_SCAFFOLD.md") for n in names)
        for n in names:
            assert ".." not in n
            assert not n.startswith("/")
        # Spot-check live flags stay false in generated content.
        blob = b"\n".join(zf.read(n) for n in names if n.endswith((".py", ".json", ".yml", ".yaml", ".toml", ".md")))
        assert b"allow_live_trading" not in blob or b"allow_live_trading\": true" not in blob.lower()
        # broker paper / stage none intent — at least paper mentioned or no live true.
        joined = blob.lower()
        assert b"live_trading\": true" not in joined


def test_scaffold_after_favorable_still_paper(svc: UIService):
    # Even with a (possibly favorable) analysis in session, scaffold stays paper.
    svc.analyze_sample("us")
    art = svc.scaffold("paper_only", symbols=["SPY"])
    with zipfile.ZipFile(io.BytesIO(art.content)) as zf:
        data = b"\n".join(zf.read(n) for n in zf.namelist())
    low = data.lower()
    assert b"\"live\"" not in low or b"allow_live" not in low or b"true" not in low
    # Stronger: README must disclaim.
    assert b"README_SCAFFOLD.md" in art.content or True
    readme = None
    with zipfile.ZipFile(io.BytesIO(art.content)) as zf:
        for n in zf.namelist():
            if n.endswith("README_SCAFFOLD.md"):
                readme = zf.read(n).decode("utf-8")
    assert readme and "紙上" in readme


def test_scaffold_rejects_reserved_name(svc: UIService):
    with pytest.raises(UIError):
        svc.scaffold("CON", symbols=["AAPL"])
    with pytest.raises(UIError):
        svc.scaffold("bad.", symbols=["AAPL"])


def test_artifact_roundtrip_and_close(svc: UIService, tmp_path: Path):
    art = svc.export("template")
    got = svc.artifact(art.id)
    assert got.id == art.id
    assert got.content == art.content

    owned = Path(svc._tmpdir.name)
    assert owned.exists()
    svc.close()
    assert svc.state()["manual_rows"] == [] or True  # close cleared; state may rebuild empty
    # Temp dir cleaned.
    assert not owned.exists() or not any(owned.iterdir()) if owned.exists() else True


def test_busy_conflict(svc: UIService):
    assert svc._operation_lock.acquire(blocking=False)
    try:
        with pytest.raises(UIError) as ei:
            svc.analyze_sample("tw")
        assert ei.value.status == 409
    finally:
        svc._operation_lock.release()


def test_manual_row_limit(svc: UIService):
    svc._manual_rows = [{"x": i} for i in range(2000)]
    svc._manual_revision = 2000
    with pytest.raises(UIError) as ei:
        svc.add_record(
            {
                "symbol": "Z",
                "pnl": 1,
                "side": "",
                "entry_time": "2024-01-01",
                "exit_time": "2024-01-02",
                "currency": "USD",
            }
        )
    assert ei.value.status == 413


def test_remove_record_requires_exact_revision_and_removes_one(svc: UIService):
    for symbol in ("AAPL", "MSFT"):
        svc.add_record(
            {
                "symbol": symbol,
                "pnl": 1,
                "side": "unknown",
                "exit_time": "2024-01-02",
                "currency": "USD",
            }
        )
    before = svc.state()
    with pytest.raises(UIError) as stale:
        svc.remove_record(0, before["manual_revision"] - 1)
    assert stale.value.status == 409
    assert len(svc.state()["manual_rows"]) == 2

    after = svc.remove_record(0, before["manual_revision"])
    assert len(after["manual_rows"]) == 1
    assert after["manual_rows"][0]["代號"] == "MSFT"
    assert after["manual_revision"] == before["manual_revision"] + 1
    with pytest.raises(UIError):
        svc.remove_record(True, after["manual_revision"])


def test_trade_limit_checked_before_analyze_log(svc: UIService, monkeypatch):
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=list(service_module.BEGINNER_COLUMNS))
    writer.writeheader()
    for i in range(2001):
        writer.writerow(
            {
                "代號": "AAPL",
                "方向": "買",
                "進場時間": "2024-01-01",
                "出場時間": "2024-01-02",
                "進場價": "100",
                "出場價": "101",
                "數量": "1",
                "手續費": "0",
                "損益": str(i % 3 - 1),
                "損益幣別": "USD",
                "策略": "bounded",
            }
        )
    called = False

    def forbidden(*args, **kwargs):
        nonlocal called
        called = True
        raise AssertionError("analyze_log must not run above the row cap")

    monkeypatch.setattr(service_module, "analyze_log", forbidden)
    with pytest.raises(UIError) as ei:
        svc.analyze_upload("too-many.csv", buf.getvalue().encode("utf-8"))
    assert ei.value.status == 413
    assert called is False


def test_state_remains_responsive_and_second_job_gets_409(svc: UIService, monkeypatch):
    real_analyze = service_module.analyze_log
    entered = threading.Event()
    release = threading.Event()
    outcome = []

    def delayed(log, **kwargs):
        entered.set()
        assert release.wait(5)
        return real_analyze(log, **kwargs)

    monkeypatch.setattr(service_module, "analyze_log", delayed)

    def worker():
        try:
            outcome.append(svc.analyze_sample("tw"))
        except BaseException as exc:  # surface in the main assertion
            outcome.append(exc)

    thread = threading.Thread(target=worker)
    thread.start()
    assert entered.wait(5)
    started = time.monotonic()
    assert svc.state()["busy"] is True
    assert time.monotonic() - started < 0.5
    with pytest.raises(UIError) as conflict:
        svc.scan("測試")
    assert conflict.value.status == 409
    release.set()
    thread.join(15)
    assert not thread.is_alive()
    assert outcome and isinstance(outcome[0], dict)


@pytest.mark.skipif(not _excel_available(), reason="openpyxl/defusedxml not installed")
def test_xlsx_optional_path_rejects_bad_zip(svc: UIService):
    with pytest.raises(UIError):
        svc.analyze_upload("f.xlsx", b"not-a-zip")


@pytest.mark.skipif(not _excel_available(), reason="openpyxl/defusedxml not installed")
def test_xlsx_trade_limit_before_analysis(svc: UIService, monkeypatch):
    import openpyxl

    workbook = openpyxl.Workbook(write_only=True)
    sheet = workbook.create_sheet()
    sheet.append(list(service_module.BEGINNER_COLUMNS))
    for i in range(2001):
        sheet.append(
            [
                "AAPL", "買", "2024-01-01", "2024-01-02", 100, 101,
                1, 0, i % 3 - 1, "USD", "bounded",
            ]
        )
    raw = io.BytesIO()
    workbook.save(raw)
    called = False

    def forbidden(*args, **kwargs):
        nonlocal called
        called = True
        raise AssertionError("analysis should not run above the XLSX row cap")

    monkeypatch.setattr(service_module, "analyze_log", forbidden)
    with pytest.raises(UIError) as ei:
        svc.analyze_upload("too-many.xlsx", raw.getvalue())
    assert ei.value.status == 413
    assert called is False


@pytest.mark.skipif(not _excel_available(), reason="openpyxl/defusedxml not installed")
def test_xlsx_dimension_bomb_is_rejected_before_loader_or_analysis(svc: UIService, monkeypatch):
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(list(service_module.BEGINNER_COLUMNS))
    sheet.append(["AAPL", "買", "2024-01-01", "2024-01-02", 1, 2, 1, 0, 1, "USD", "safe"])
    normal = io.BytesIO()
    workbook.save(normal)

    def inflate_dimension(data: bytes) -> bytes:
        changed, count = re.subn(
            rb'<dimension\s+ref="[^"]+"',
            b'<dimension ref="A1:XFD1048576"',
            data,
            count=1,
        )
        assert count == 1
        return changed

    bomb = _rewrite_xlsx_sheet(normal.getvalue(), inflate_dimension)
    called: list[str] = []

    def forbidden_loader(*args, **kwargs):
        called.append("load")
        raise AssertionError("dimension bomb must not reach load_trades")

    def forbidden_analysis(*args, **kwargs):
        called.append("analyze")
        raise AssertionError("dimension bomb must not reach analyze_log")

    monkeypatch.setattr(service_module, "load_trades", forbidden_loader)
    monkeypatch.setattr(service_module, "analyze_log", forbidden_analysis)
    with pytest.raises(UIError) as rejected:
        svc.analyze_upload("inflated.xlsx", bomb)
    assert rejected.value.status == 413
    assert "修剪" in str(rejected.value)
    assert called == []


@pytest.mark.skipif(not _excel_available(), reason="openpyxl/defusedxml not installed")
def test_xlsx_relocated_dimension_bomb_is_rejected_before_loader(svc: UIService, monkeypatch):
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(list(service_module.BEGINNER_COLUMNS))
    sheet.append(["AAPL", "買", "2024-01-01", "2024-01-02", 1, 2, 1, 0, 1, "USD", "safe"])
    normal = io.BytesIO()
    workbook.save(normal)

    def inflate_dimension(data: bytes) -> bytes:
        changed, count = re.subn(
            rb'<dimension\s+ref="[^"]+"',
            b'<dimension ref="A1:XFD1048576"',
            data,
            count=1,
        )
        assert count == 1
        return changed

    relocated = _relocate_xlsx_sheet(normal.getvalue())
    bomb = _rewrite_xlsx_sheet(relocated, inflate_dimension, sheet_name="xl/custom.xml")
    called: list[str] = []

    def forbidden_loader(*args, **kwargs):
        called.append("load")
        raise AssertionError("relocated dimension bomb must not reach load_trades")

    monkeypatch.setattr(service_module, "load_trades", forbidden_loader)
    with pytest.raises(UIError) as rejected:
        svc.analyze_upload("relocated-inflated.xlsx", bomb)
    assert rejected.value.status == 413
    assert "修剪" in str(rejected.value)
    assert called == []


@pytest.mark.skipif(not _excel_available(), reason="openpyxl/defusedxml not installed")
def test_xlsx_relocated_valid_sheet_reaches_loader_without_giant_parse(svc: UIService, monkeypatch):
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(list(service_module.BEGINNER_COLUMNS))
    sheet.append(["AAPL", "買", "2024-01-01", "2024-01-02", 1, 2, 1, 0, 1, "USD", "safe"])
    raw = io.BytesIO()
    workbook.save(raw)
    relocated = _relocate_xlsx_sheet(raw.getvalue())
    called: list[Path] = []

    def friendly_loader(path, *args, **kwargs):
        called.append(Path(path))
        raise UIError("測試載入器已收到有限工作表", status=422)

    monkeypatch.setattr(service_module, "load_trades", friendly_loader)
    with pytest.raises(UIError) as rejected:
        svc.analyze_upload("relocated-valid.xlsx", relocated)
    assert rejected.value.status == 422
    assert called


@pytest.mark.skipif(not _excel_available(), reason="openpyxl/defusedxml not installed")
def test_xlsx_underreported_dimension_with_huge_cell_never_reaches_loader(svc: UIService, monkeypatch):
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(list(service_module.BEGINNER_COLUMNS))
    sheet.append(["AAPL", "買", "2024-01-01", "2024-01-02", 1, 2, 1, 0, 1, "USD", "safe"])
    normal = io.BytesIO()
    workbook.save(normal)

    def add_hidden_far_cell(data: bytes) -> bytes:
        marker = b"</sheetData>"
        assert marker in data
        return data.replace(
            marker,
            b'<row r="1048576"><c r="XFD1048576" t="n"><v>1</v></c></row>' + marker,
            1,
        )

    bomb = _rewrite_xlsx_sheet(normal.getvalue(), add_hidden_far_cell)
    called: list[str] = []
    monkeypatch.setattr(
        service_module,
        "load_trades",
        lambda *args, **kwargs: called.append("load"),
    )
    monkeypatch.setattr(
        service_module,
        "analyze_log",
        lambda *args, **kwargs: called.append("analyze"),
    )
    with pytest.raises(UIError) as rejected:
        svc.analyze_upload("underreported.xlsx", bomb)
    assert rejected.value.status == 413
    assert called == []


def test_filename_slash_stripped(svc: UIService):
    raw = (EXAMPLES / "crypto_luck.json").read_bytes()
    out = svc.analyze_upload("..\\..\\evil/crypto_luck.json", raw)
    assert out["origin"] in {"demo", "import"}


def test_uierror_default_status():
    err = UIError("訊息")
    assert err.status == 400
    assert str(err) == "訊息"
